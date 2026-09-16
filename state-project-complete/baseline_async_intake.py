"""Fast-acknowledge Evidence intake for long Baseline Setup sources.

Baseline dogfooding showed a long source could hold the browser open for around
80 seconds even when the eventual interpretation succeeded. This module keeps
Evidence capture and AI interpretation separate for Baseline Setup:

1. save the original Evidence atomically and return immediately;
2. analyze it in a FastAPI background task;
3. let the existing Baseline draft/coverage APIs expose pending/failed status.

Nothing here authorizes Current State. The normal interpretation pipeline still
creates draft proposals/Reviews, and a person still confirms Starting State.
"""
from __future__ import annotations

import io
import logging
from typing import Any

from docx import Document as DocxDocument
from fastapi import BackgroundTasks, File, HTTPException, Request, UploadFile
from pydantic import BaseModel, ConfigDict, Field, field_validator
from pypdf import PdfReader

from baseline_setup import is_baseline_setup
from db import connect
from interpretation_pipeline_integrated import new_id, process_evidence

logger = logging.getLogger("state.baseline.intake")

_UPLOAD_EXTENSIONS = {".txt", ".md", ".pdf", ".docx"}
_UPLOAD_MAX_BYTES = 15_000_000


class BaselineEvidenceInput(BaseModel):
    model_config = ConfigDict(extra="forbid")
    content: str = Field(min_length=1, max_length=100_000)
    source_type: str = Field(default="manual_note", min_length=1, max_length=80)

    @field_validator("content")
    @classmethod
    def content_not_blank(cls, value: str) -> str:
        value = value.strip()
        if not value:
            raise ValueError("content must not be blank")
        return value


def _unreadable(message: str, code: str = "unreadable_file") -> HTTPException:
    return HTTPException(
        status_code=422,
        detail={"code": code, "error_details": {"error_message": message}},
    )


def _extract_upload_text(raw: bytes, extension: str) -> str:
    """Mirror the existing State upload contract for Baseline fast intake."""
    if extension in (".txt", ".md"):
        try:
            return raw.decode("utf-8")
        except UnicodeDecodeError as exc:
            raise _unreadable("That file doesn't look like plain text State can read.") from exc

    if extension == ".pdf":
        try:
            reader = PdfReader(io.BytesIO(raw))
            if reader.is_encrypted:
                raise _unreadable("That PDF is password-protected. Remove the password and try again.")
            pages_text = [page.extract_text() or "" for page in reader.pages]
        except HTTPException:
            raise
        except Exception as exc:
            raise _unreadable("That PDF could not be read. It may be corrupted.") from exc
        text = "\n\n".join(pages_text).strip()
        if not text:
            raise _unreadable(
                "State couldn't find readable text in that PDF. Scanned or image-only PDFs aren't supported yet -- try a text-based PDF instead.",
                code="no_extractable_text",
            )
        return text

    if extension == ".docx":
        try:
            document = DocxDocument(io.BytesIO(raw))
        except Exception as exc:
            raise _unreadable("That .docx file could not be read. It may be corrupted.") from exc
        parts = [paragraph.text for paragraph in document.paragraphs if paragraph.text.strip()]
        for table in document.tables:
            for row in table.rows:
                row_text = " | ".join(cell.text.strip() for cell in row.cells if cell.text.strip())
                if row_text:
                    parts.append(row_text)
        return "\n\n".join(parts)

    raise _unreadable("Unsupported file type.", code="unsupported_file_type")


def _selected_provider(app: Any, settings: Any):
    provider = app.state.provider
    if provider is not None:
        return provider
    # api_core already owns provider/env construction. Reuse it instead of
    # inventing a second provider configuration path.
    from api_core import _provider_from_env

    provider = _provider_from_env(settings)
    app.state.provider = provider
    return provider


def _project_id_from_request(request: Request) -> str:
    project_id = (request.headers.get("X-State-Project-Id") or "").strip()
    if not project_id:
        raise HTTPException(status_code=400, detail="X-State-Project-Id is required")
    return project_id


def _store_pending_evidence(settings: Any, project_id: str, content: str, source_type: str, source_name: str | None) -> str:
    connection = connect(settings.connection_url())
    try:
        connection.project_id = project_id
        exists = connection.execute("SELECT id FROM projects WHERE id=?", (project_id,)).fetchone()
        if exists is None:
            raise HTTPException(status_code=404, detail="Project not found")
        if not is_baseline_setup(connection):
            raise HTTPException(status_code=409, detail={"code": "baseline_not_active"})

        evidence_id = new_id("evidence")
        connection.execute(
            "INSERT INTO evidence(id, content, source_type, source_name, project_id) VALUES (?, ?, ?, ?, ?)",
            (evidence_id, content, source_type, source_name, project_id),
        )
        connection.commit()
        return evidence_id
    finally:
        connection.close()


def _analyze_in_background(settings: Any, provider: Any, project_id: str, evidence_id: str) -> None:
    connection = connect(settings.connection_url())
    try:
        connection.project_id = project_id
        process_evidence(connection, evidence_id=evidence_id, provider=provider)
    except Exception:
        logger.exception("Background Baseline analysis failed for %s", evidence_id)
        # Most provider/contract failures are already persisted by
        # process_evidence. This fallback handles unexpected failures so the
        # UI does not display an Evidence item as "Analyzing" forever.
        try:
            connection.rollback()
        except Exception:
            pass
        try:
            connection.execute(
                "UPDATE evidence SET processing_status='failed' WHERE id=? AND project_id=? AND processing_status='pending'",
                (evidence_id, project_id),
            )
            connection.commit()
        except Exception:
            logger.exception("Could not mark failed Baseline Evidence %s", evidence_id)
    finally:
        connection.close()


def _queue(background_tasks: BackgroundTasks, settings: Any, provider: Any, project_id: str, evidence_id: str) -> dict:
    background_tasks.add_task(_analyze_in_background, settings, provider, project_id, evidence_id)
    return {
        "evidence_id": evidence_id,
        "interpretation_record_id": None,
        "processing_status": "pending",
        "reviews": [],
    }


def register_baseline_async_intake_routes(app: Any, settings: Any) -> None:
    @app.post("/api/baseline/evidence", status_code=202)
    def post_baseline_evidence(
        payload: BaselineEvidenceInput,
        request: Request,
        background_tasks: BackgroundTasks,
    ) -> dict:
        project_id = _project_id_from_request(request)
        provider = _selected_provider(request.app, settings)
        evidence_id = _store_pending_evidence(
            settings,
            project_id,
            payload.content.strip(),
            payload.source_type,
            None,
        )
        return _queue(background_tasks, settings, provider, project_id, evidence_id)

    @app.post("/api/baseline/evidence/upload", status_code=202)
    async def upload_baseline_evidence(
        request: Request,
        background_tasks: BackgroundTasks,
        file: UploadFile = File(...),
    ) -> dict:
        filename = (file.filename or "").strip()
        extension = filename[filename.rfind("."):].lower() if "." in filename else ""
        if extension not in _UPLOAD_EXTENSIONS:
            raise HTTPException(
                status_code=422,
                detail={"code": "unsupported_file_type", "error_details": {
                    "error_message": "State can currently read .txt, .md, .pdf, and .docx files. Other formats aren't supported yet.",
                }},
            )
        raw = await file.read()
        if len(raw) > _UPLOAD_MAX_BYTES:
            raise HTTPException(
                status_code=422,
                detail={"code": "file_too_large", "error_details": {
                    "error_message": "That file is too large to upload as Evidence right now.",
                }},
            )
        content = _extract_upload_text(raw, extension).strip()
        if not content:
            raise HTTPException(
                status_code=422,
                detail={"code": "empty_file", "error_details": {
                    "error_message": "That file has no readable text to add as Evidence.",
                }},
            )
        if len(content) > 100_000:
            content = content[:100_000]

        project_id = _project_id_from_request(request)
        provider = _selected_provider(request.app, settings)
        evidence_id = _store_pending_evidence(
            settings,
            project_id,
            content,
            "uploaded_note",
            filename or None,
        )
        return _queue(background_tasks, settings, provider, project_id, evidence_id)
